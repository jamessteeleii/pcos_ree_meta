"""Build the Obesity Reviews submission package without changing the preprint.

The script derives two Word manuscripts from submission_manuscript.qmd:

* a clean version with all revised wording accepted;
* a highlighted version that preserves the source's red revision markup and
  encodes those red runs as Word insertions.

Both manuscripts place the table and all legends after the references. Figure
artwork is supplied separately as 300 dpi TIFF files.
"""

from __future__ import annotations

import argparse
import copy
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from PIL import Image


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}
REVISION_AUTHOR = "James Steele"
REVISION_DATE = "2026-10-05T00:00:00Z"


def split_front_matter(text: str) -> tuple[str, str]:
    match = re.match(r"\A---\r?\n(.*?)\r?\n---\r?\n", text, flags=re.S)
    if not match:
        raise ValueError("Could not locate Quarto front matter.")
    return match.group(1), text[match.end() :]


def make_docx_front_matter(front: str) -> str:
    front = re.sub(
        r"(?ms)^format:\r?\n.*?(?=^geometry:)",
        "format:\n  docx:\n    reference-doc: submission_reference.docx\n",
        front,
    )
    front = re.sub(r"(?m)^geometry:\r?\n(?:  - [^\r\n]*\r?\n)+", "", front)
    return front


def find_matching_brace(text: str, open_index: int) -> int:
    depth = 0
    escaped = False
    for index in range(open_index, len(text)):
        char = text[index]
        if escaped:
            escaped = False
            continue
        if char == "\\":
            escaped = True
            continue
        if char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                return index
    raise ValueError("Unbalanced braces in revision markup.")


def replace_textcolor(text: str, tracked: bool) -> str:
    marker = r"\textcolor{red}{"
    while marker in text:
        start = text.rfind(marker)
        open_index = start + len(marker) - 1
        end = find_matching_brace(text, open_index)
        content = text[open_index + 1 : end]
        replacement = f'[{content}]{{custom-style="Revision"}}' if tracked else content
        text = text[:start] + replacement + text[end + 1 :]
    return text


def latex_links_to_markdown(text: str) -> str:
    marker = r"\href{"
    while marker in text:
        start = text.rfind(marker)
        url_open = start + len(marker) - 1
        url_end = find_matching_brace(text, url_open)
        label_open = url_end + 1
        if label_open >= len(text) or text[label_open] != "{":
            raise ValueError("Malformed LaTeX hyperlink.")
        label_end = find_matching_brace(text, label_open)
        url = text[url_open + 1 : url_end]
        label = text[label_open + 1 : label_end]
        text = text[:start] + f"[{label}]({url})" + text[label_end + 1 :]
    return text


def mark_red_groups(body: str, tracked: bool) -> str:
    lines = body.splitlines()
    output: list[str] = []
    in_group = False
    in_code = False
    pending_begin = False

    for line in lines:
        stripped = line.strip()
        if stripped == r"\begingroup":
            pending_begin = True
            continue
        if pending_begin and stripped == r"\color{red}":
            in_group = True
            pending_begin = False
            continue
        if pending_begin:
            output.append(r"\begingroup")
            pending_begin = False

        if stripped == r"\endgroup" and in_group:
            in_group = False
            continue

        if stripped.startswith("```"):
            in_code = not in_code
            output.append(line)
            continue

        if tracked and in_group and stripped and not in_code:
            indent = line[: len(line) - len(line.lstrip())]
            content = line.lstrip()
            heading = re.match(r"^(#{1,6})\s+(.*)$", content)
            if heading:
                line = (
                    f'{indent}{heading.group(1)} '
                    f'[{heading.group(2)}]{{custom-style="Revision"}}'
                )
            elif content.startswith(("- ", "+ ", "* ")):
                line = (
                    f'{indent}{content[:2]}'
                    f'[{content[2:]}]{{custom-style="Revision"}}'
                )
            else:
                line = f'{indent}[{content}]{{custom-style="Revision"}}'
        output.append(line)

    if pending_begin:
        output.append(r"\begingroup")
    if in_group:
        raise ValueError("Unclosed red revision group.")
    return "\n".join(output) + "\n"


def remove_figures(body: str) -> str:
    body = re.sub(
        r"(?ms)^\\begin\{figure\}\[htbp\].*?^\\end\{figure\}\s*",
        "",
        body,
    )
    body = re.sub(
        r"(?ms)^```\{r\}\r?\n#\| label: fig-mean-REE.*?^```\s*",
        "",
        body,
    )
    body = re.sub(
        r"(?ms)^```\{r\}\r?\n#\| label: fig-variance-REE.*?^```\s*",
        "",
        body,
    )
    return body


def extract_table_chunk(body: str) -> tuple[str, str]:
    start = body.find("analysis_summary <- tibble(")
    if start < 0:
        raise ValueError("Could not locate analysis summary table code.")
    fence_start = body.rfind("```{r}", 0, start)
    fence_end = body.find("```", start)
    if fence_start < 0 or fence_end < 0:
        raise ValueError("Could not locate analysis summary code fences.")
    fence_end += 3
    chunk = body[fence_start:fence_end]
    chunk = re.sub(
        r"caption\s*=\s*\"Summary of primary, sensitivity, and exploratory analyses\.\"",
        "caption = NULL",
        chunk,
    )
    body = body[:fence_start] + body[fence_end:]
    return body, chunk


def normalize_latex_for_docx(body: str) -> str:
    body = body.replace(r"Figure~\ref{fig-prisma}", "Figure 1")
    body = body.replace(r"Figure \ref{fig-prisma}", "Figure 1")
    body = body.replace(r"Figure~\ref{fig-mean-REE}", "Figure 2")
    body = body.replace(r"Figure~\ref{fig-variance-REE}", "Figure 3")
    body = body.replace(r"\newpage", "{{< pagebreak >}}")
    body = latex_links_to_markdown(body)
    body = body.replace("~", " ")
    return body


def legends_and_table(table_chunk: str, tracked: bool) -> str:
    entries = [
        (
            "Table 1: Summary of primary, sensitivity, and exploratory analyses. "
            "Values are posterior point estimates with 95% quantile intervals (QI). "
            "Variability is expressed as a PMOS-to-control standard-deviation ratio, "
            "except for the direct pairwise model, which uses a coefficient-of-variation ratio."
        ),
        (
            "Figure 1: Updated PRISMA 2020 flow diagram for the systematic review. "
            "The updated database search covered January 2025 through September 2026 and "
            "identified no additional eligible studies. The website shown in the original-search "
            "pathway was a prior narrative review by members of the author group that identified "
            "two studies not found in the systematic database search."
        ),
        (
            "Figure 2: Posterior distributions, point estimates, and 95% quantile intervals for "
            "the study-level conditional estimates, global estimates by condition, and the "
            "contrast between conditions for mean resting energy expenditure."
        ),
        (
            "Figure 3: Posterior distributions, point estimates, and 95% quantile intervals for "
            "the study-level conditional estimates, global estimates by condition, and the "
            "contrast ratio between conditions for the between-person standard deviation of "
            "resting energy expenditure."
        ),
    ]
    if tracked:
        entries = [f'[{entry}]{{custom-style="Revision"}}' for entry in entries]
        table_title = (
            '[Table 1. Summary of primary, sensitivity, and exploratory analyses.]'
            '{custom-style="Revision"}'
        )
        table_note = (
            '[Note. Values are posterior point estimates with 95% quantile intervals (QI). '
            'Variability is expressed as a PMOS-to-control standard-deviation ratio, except for '
            'the direct pairwise model, which uses a coefficient-of-variation ratio.]'
            '{custom-style="Revision"}'
        )
    else:
        table_title = "Table 1. Summary of primary, sensitivity, and exploratory analyses."
        table_note = (
            "Note. Values are posterior point estimates with 95% quantile intervals (QI). "
            "Variability is expressed as a PMOS-to-control standard-deviation ratio, except for "
            "the direct pairwise model, which uses a coefficient-of-variation ratio."
        )

    return (
        "\n{{< pagebreak >}}\n\n# Table and Figure Legends\n\n"
        + "\n\n".join(entries)
        + "\n\n{{< pagebreak >}}\n\n"
        + table_title
        + "\n\n"
        + table_chunk
        + "\n\n"
        + table_note
        + "\n"
    )


def build_qmd(source: str, tracked: bool) -> str:
    front, body = split_front_matter(source)
    front = make_docx_front_matter(front)
    body = re.sub(
        r'(?ms)pre_print_project_root <- if \(file\.exists\(file\.path\("manuscript", "pre_print_objects\.rds"\)\)\) \{\r?\n  "\."\r?\n\} else \{\r?\n  "\.\."\r?\n\}',
        'pre_print_project_root <- normalizePath(Sys.getenv("PCOS_PROJECT_ROOT"), mustWork = TRUE)',
        body,
    )
    body = remove_figures(body)
    body, table_chunk = extract_table_chunk(body)
    body = mark_red_groups(body, tracked=tracked)
    body = latex_links_to_markdown(body)
    body = replace_textcolor(body, tracked=tracked)
    body = normalize_latex_for_docx(body)
    body = body.replace(
        "None of the additional sensitivity or exploratory analyses materially altered the conclusion. In the table, variability",
        "The primary, sensitivity, and exploratory estimates are summarised in Table 1. None of the additional sensitivity or exploratory analyses materially altered the conclusion. In Table 1, variability",
    )
    body = re.sub(
        r"(?m)^# References\s*$",
        "# References\n\n::: {#refs}\n:::",
        body,
        count=1,
    )
    body = body.rstrip() + legends_and_table(table_chunk, tracked=tracked)
    return f"---\n{front}\n---\n{body}"


def set_font(run, name: str, size: float | None = None, color: RGBColor | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color


def add_style_if_missing(document: Document, name: str, style_type: WD_STYLE_TYPE):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, style_type)


def create_reference_docx(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)

    for name, size, bold in (
        ("Title", 14, True),
        ("Heading 1", 12, True),
        ("Heading 2", 12, True),
        ("Heading 3", 12, True),
    ):
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True

    revision = add_style_if_missing(document, "Revision", WD_STYLE_TYPE.CHARACTER)
    revision.font.name = "Times New Roman"
    revision.font.size = Pt(12)
    revision.font.color.rgb = RGBColor(255, 0, 0)
    revision._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    revision._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")

    document.save(path)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_font(run, "Times New Roman", 10)


def insert_paragraph_after(paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    new_paragraph = Paragraph(new_element, paragraph._parent)
    if style:
        new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def remove_paragraph_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_title_page_affiliations(document: Document) -> None:
    author_affiliations = {
        "Richard Kirwan": "1",
        "Leigh Peele": "2",
        "Gregory Nuckols": "2",
        "Georgia Kohlhoff": "3",
        "Hannah Cabré": "4",
        "Alyssa Olenick": "5",
        "James Steele": "2, 6, 7",
    }
    affiliations = [
        "1 Research Institute for Sports and Exercise Sciences, Liverpool John Moores University, Liverpool, UK",
        "2 MacroFactor, Stronger by Science Technologies LLC, Raleigh, North Carolina, USA",
        "3 Flourishing Health, Liverpool, UK",
        "4 Reproductive Endocrinology and Women’s Health Research Program, Pennington Biomedical Research Center, Baton Rouge, Louisiana, USA",
        "5 Dr Alyssa Olenick LLC",
        "6 Steele Research Limited, Eastleigh, Hampshire, UK",
        "7 School of Health and Biomedical Sciences, Royal Melbourne Institute of Technology, Melbourne, VIC, Australia",
    ]

    by_text = {paragraph.text.strip(): paragraph for paragraph in document.paragraphs}
    for author, affiliation_ids in author_affiliations.items():
        paragraph = by_text.get(author)
        if paragraph is None:
            raise ValueError(f"Could not locate title-page author: {author}")
        run = paragraph.add_run(affiliation_ids)
        run.font.superscript = True
        set_font(run, "Times New Roman", 9)

    anchor = by_text["James Steele"]
    for affiliation in affiliations:
        anchor = insert_paragraph_after(anchor, affiliation, style="Normal")
        anchor.paragraph_format.line_spacing = 1
        anchor.paragraph_format.space_after = Pt(0)
        for run in anchor.runs:
            set_font(run, "Times New Roman", 10)

    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "Financial Disclosures/Conflicts of Interest":
            break
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.space_after = Pt(0)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def iter_paragraphs(document: Document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                yield paragraph


def make_journal_layout(path: Path, tracked: bool) -> None:
    document = Document(path)
    add_title_page_affiliations(document)
    footnote_reference = add_style_if_missing(
        document, "Footnote Reference", WD_STYLE_TYPE.CHARACTER
    )
    footnote_reference.font.name = "Times New Roman"
    footnote_reference.font.size = Pt(8)
    footnote_reference.font.superscript = True
    footnote_reference._element.get_or_add_rPr().rFonts.set(
        qn("w:ascii"), "Times New Roman"
    )
    footnote_reference._element.get_or_add_rPr().rFonts.set(
        qn("w:hAnsi"), "Times New Roman"
    )
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)
        sect_pr = section._sectPr
        line_numbering = sect_pr.find(qn("w:lnNumType"))
        if line_numbering is None:
            line_numbering = OxmlElement("w:lnNumType")
            sect_pr.append(line_numbering)
        line_numbering.set(qn("w:countBy"), "1")
        line_numbering.set(qn("w:restart"), "newPage")
        footer = section.footer
        footer.is_linked_to_previous = False
        for paragraph in footer.paragraphs:
            paragraph._element.getparent().remove(paragraph._element)
        add_page_field(footer.add_paragraph())

    for style_name in ("Normal", "Body Text", "First Paragraph"):
        if style_name in document.styles:
            style = document.styles[style_name]
            style.font.name = "Times New Roman"
            style.font.size = Pt(12)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
            style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
            style.paragraph_format.line_spacing = 2
            style.paragraph_format.space_after = Pt(0)

    for heading_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        if heading_name in document.styles:
            style = document.styles[heading_name]
            style.font.name = "Times New Roman"
            style.font.color.rgb = RGBColor(0, 0, 0)
            style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
            style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
            p_pr = style._element.get_or_add_pPr()
            border = p_pr.find(qn("w:pBdr"))
            if border is not None:
                p_pr.remove(border)

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if paragraph.style is not None and paragraph.style.name == "Title":
            remove_paragraph_border(paragraph)
        paragraph.paragraph_format.line_spacing = 2
        paragraph.paragraph_format.space_after = Pt(0)
        if text == "Table and Figure Legends":
            paragraph.paragraph_format.page_break_before = True
        if text.startswith("Table 1. Summary of primary, sensitivity, and exploratory analyses"):
            paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            is_revision = run.style is not None and run.style.name == "Revision"
            color = RGBColor(255, 0, 0) if tracked and is_revision else None
            set_font(run, "Times New Roman", 12, color=color)

    table_title_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.strip().startswith(
            "Table 1. Summary of primary, sensitivity, and exploratory analyses"
        )
    ]
    if len(table_title_paragraphs) >= 2:
        table_title_paragraphs[-1].paragraph_format.page_break_before = True
        table_title_paragraphs[-1].paragraph_format.keep_with_next = True

    for table in document.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table)
        widths = [Inches(1.55), Inches(1.05), Inches(2.15), Inches(1.75)]
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                if column_index < len(widths):
                    cell.width = widths[column_index]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell)
                if row_index == 0:
                    shade_cell(cell, "D9E2F3")
                elif row_index % 2 == 0:
                    shade_cell(cell, "F7F9FC")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.LEFT
                        if column_index == 0
                        else WD_ALIGN_PARAGRAPH.CENTER
                    )
                    for run in paragraph.runs:
                        set_font(
                            run,
                            "Times New Roman",
                            9.5,
                            color=RGBColor(255, 0, 0) if tracked else None,
                        )
                        if row_index == 0:
                            run.bold = True
        header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_tr_pr.append(repeat)

    document.core_properties.title = (
        "Resting energy expenditure of women with and without polyendocrine metabolic ovarian syndrome"
    )
    document.core_properties.subject = (
        "Obesity Reviews submission clean manuscript"
        if not tracked
        else "Obesity Reviews submission highlighted manuscript"
    )
    document.core_properties.author = "James Steele and co-authors"
    document.core_properties.last_modified_by = "James Steele and co-authors"
    document.save(path)


def format_footnote_markers(path: Path) -> None:
    """Use lower-case letters and true superscript formatting for footnotes."""
    from lxml import etree

    temporary = path.with_suffix(".footnotes.tmp.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", zipfile.ZIP_DEFLATED
    ) as destination:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename in {"word/document.xml", "word/footnotes.xml"}:
                root = etree.fromstring(data)
                for run in root.xpath(
                    ".//w:r[w:footnoteReference or w:footnoteRef]", namespaces=NS
                ):
                    run_properties = run.find(qn("w:rPr"))
                    if run_properties is None:
                        run_properties = etree.Element(qn("w:rPr"))
                        run.insert(0, run_properties)

                    style = run_properties.find(qn("w:rStyle"))
                    if style is None:
                        style = etree.Element(qn("w:rStyle"))
                        run_properties.insert(0, style)
                    style.set(qn("w:val"), "FootnoteReference")

                    for tag, value in (
                        ("w:sz", "16"),
                        ("w:szCs", "16"),
                        ("w:vertAlign", "superscript"),
                    ):
                        element = run_properties.find(qn(tag))
                        if element is None:
                            element = etree.Element(qn(tag))
                            run_properties.append(element)
                        element.set(
                            qn("w:val"), "superscript" if tag == "w:vertAlign" else value
                        )

                if item.filename == "word/document.xml":
                    for section_properties in root.xpath(".//w:sectPr", namespaces=NS):
                        footnote_properties = section_properties.find(qn("w:footnotePr"))
                        if footnote_properties is None:
                            footnote_properties = etree.Element(qn("w:footnotePr"))
                            section_properties.insert(0, footnote_properties)
                        number_format = footnote_properties.find(qn("w:numFmt"))
                        if number_format is None:
                            number_format = etree.Element(qn("w:numFmt"))
                            footnote_properties.insert(0, number_format)
                        number_format.set(qn("w:val"), "lowerLetter")

                data = etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone="yes"
                )
            destination.writestr(item, data)
    temporary.replace(path)


def revision_run(run) -> bool:
    r_pr = run.find(qn("w:rPr"))
    if r_pr is None:
        return False
    style = r_pr.find(qn("w:rStyle"))
    if style is not None and style.get(qn("w:val")) == "Revision":
        return True
    color = r_pr.find(qn("w:color"))
    if color is not None and color.get(qn("w:val"), "").upper() == "FF0000":
        return True
    return False


def add_true_insertions(path: Path) -> int:
    temporary = path.with_suffix(".tracked.tmp.docx")
    count = 0
    next_id = 1
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", zipfile.ZIP_DEFLATED
    ) as destination:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/settings.xml":
                from lxml import etree

                root = etree.fromstring(data)
                if root.find(qn("w:trackRevisions")) is None:
                    track = etree.Element(qn("w:trackRevisions"))
                    root.insert(0, track)
                data = etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone="yes"
                )
            elif item.filename in {"word/document.xml", "word/footnotes.xml"}:
                from lxml import etree

                root = etree.fromstring(data)
                runs = list(root.xpath(".//w:r", namespaces=NS))
                for run in runs:
                    if not revision_run(run):
                        continue
                    parent = run.getparent()
                    if parent.tag == qn("w:ins"):
                        continue
                    insertion = etree.Element(qn("w:ins"))
                    insertion.set(qn("w:id"), str(next_id))
                    insertion.set(qn("w:author"), REVISION_AUTHOR)
                    insertion.set(qn("w:date"), REVISION_DATE)
                    parent.replace(run, insertion)
                    insertion.append(run)
                    next_id += 1
                    count += 1
                data = etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone="yes"
                )
            destination.writestr(item, data)
    temporary.replace(path)
    return count


def report_revisions(path: Path) -> dict[str, int]:
    from lxml import etree

    totals = {"insertions": 0, "deletions": 0, "red_runs": 0, "images": 0}
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("word/media/"):
                totals["images"] += 1
            if name not in {"word/document.xml", "word/footnotes.xml"}:
                continue
            root = etree.fromstring(archive.read(name))
            totals["insertions"] += len(root.xpath(".//w:ins", namespaces=NS))
            totals["deletions"] += len(root.xpath(".//w:del", namespaces=NS))
            totals["red_runs"] += sum(
                1 for run in root.xpath(".//w:r", namespaces=NS) if revision_run(run)
            )
    return totals


def copy_figures(project_root: Path, output_dir: Path) -> None:
    prisma_source = project_root / "manuscript" / "prisma.png"
    with Image.open(prisma_source) as image:
        converted = image.convert("RGB")
        converted.save(
            output_dir / "Figure_1_PRISMA.tiff",
            format="TIFF",
            compression="tiff_lzw",
            dpi=(300, 300),
        )
    shutil.copy2(
        project_root / "plots" / "combine_mean_plot.tiff",
        output_dir / "Figure_2_Mean_REE.tiff",
    )
    shutil.copy2(
        project_root / "plots" / "combine_variance_plot.tiff",
        output_dir / "Figure_3_REE_variability.tiff",
    )


def validate_figures(output_dir: Path) -> None:
    expected = {
        "Figure_1_PRISMA.tiff": (4961, 3508),
        "Figure_2_Mean_REE.tiff": (3000, 1500),
        "Figure_3_REE_variability.tiff": (3000, 1500),
    }
    for filename, dimensions in expected.items():
        path = output_dir / filename
        with Image.open(path) as image:
            if image.format != "TIFF" or image.size != dimensions:
                raise ValueError(f"Unexpected TIFF properties for {path}: {image.format}, {image.size}")
            dpi = image.info.get("dpi", (0, 0))
            if min(dpi) < 299:
                raise ValueError(f"Figure resolution is below 300 dpi: {path} ({dpi})")


def run_quarto(qmd_path: Path, project_root: Path, r_bin: Path) -> Path:
    quarto = shutil.which("quarto")
    if not quarto:
        raise RuntimeError("Quarto was not found on PATH.")
    env = os.environ.copy()
    env["PCOS_PROJECT_ROOT"] = str(project_root)
    env["R_PROFILE_USER"] = "NUL"
    env["R_ENVIRON_USER"] = "NUL"
    env["PATH"] = str(r_bin) + os.pathsep + env.get("PATH", "")
    subprocess.run(
        [quarto, "render", qmd_path.name, "--to", "docx"],
        cwd=qmd_path.parent,
        env=env,
        check=True,
    )
    output = qmd_path.with_suffix(".docx")
    if not output.exists():
        raise FileNotFoundError(output)
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--r-bin", type=Path, required=True)
    args = parser.parse_args()

    project_root = args.project_root.resolve()
    output_dir = (
        args.output_dir.resolve()
        if args.output_dir
        else project_root / "manuscript" / "submission_package"
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    source_path = project_root / "manuscript" / "submission_manuscript.qmd"
    source = source_path.read_text(encoding="utf-8")

    with tempfile.TemporaryDirectory(prefix="pcos_submission_") as temporary_directory:
        build_dir = Path(temporary_directory)
        shutil.copy2(project_root / "manuscript" / "references.bib", build_dir / "references.bib")
        shutil.copy2(
            project_root / "manuscript" / "american-medical-association.csl",
            build_dir / "american-medical-association.csl",
        )
        reference_docx = build_dir / "submission_reference.docx"
        create_reference_docx(reference_docx)

        clean_qmd = build_dir / "submission_clean.qmd"
        tracked_qmd = build_dir / "submission_tracked_changes.qmd"
        clean_qmd.write_text(build_qmd(source, tracked=False), encoding="utf-8")
        tracked_qmd.write_text(build_qmd(source, tracked=True), encoding="utf-8")

        clean_render = run_quarto(clean_qmd, project_root, args.r_bin)
        tracked_render = run_quarto(tracked_qmd, project_root, args.r_bin)

        clean_output = output_dir / "Manuscript_Clean.docx"
        tracked_output = output_dir / "Manuscript_Tracked_Changes.docx"
        shutil.copy2(clean_render, clean_output)
        shutil.copy2(tracked_render, tracked_output)
        make_journal_layout(clean_output, tracked=False)
        make_journal_layout(tracked_output, tracked=True)
        format_footnote_markers(clean_output)
        format_footnote_markers(tracked_output)
        insertion_count = add_true_insertions(tracked_output)
        if insertion_count == 0:
            raise ValueError("The tracked manuscript contains no Word insertion markup.")

    copy_figures(project_root, output_dir)
    validate_figures(output_dir)

    clean_report = report_revisions(output_dir / "Manuscript_Clean.docx")
    tracked_report = report_revisions(output_dir / "Manuscript_Tracked_Changes.docx")
    if clean_report["insertions"] or clean_report["deletions"] or clean_report["red_runs"]:
        raise ValueError(f"Clean manuscript still contains revision markup: {clean_report}")
    if clean_report["images"] or tracked_report["images"]:
        raise ValueError("Figure artwork must not be embedded in the main manuscripts.")
    if tracked_report["insertions"] == 0 or tracked_report["red_runs"] == 0:
        raise ValueError(f"Tracked manuscript lacks highlighted insertions: {tracked_report}")

    print(f"Clean manuscript: {clean_report}")
    print(f"Tracked manuscript: {tracked_report}")
    print(f"Output directory: {output_dir}")


if __name__ == "__main__":
    main()
